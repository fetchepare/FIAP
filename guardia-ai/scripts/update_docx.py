"""Script para preencher os itens 'XXXXX' no arquivo 'Tech Challenge 05.docx'."""
import os
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_document(docx_path: str):
    print(f"Processando arquivo: {docx_path}")
    doc = docx.Document(docx_path)
    
    # 1. P[32] - DESAFIO
    doc.paragraphs[32].text = (
        "Guardiã AI — Inteligência Artificial para Saúde e Segurança da Mulher\n\n"
        "Uma instituição que atua no acolhimento e assistência à mulher busca aprimorar seus fluxos "
        "de atendimento utilizando Inteligência Artificial. Durante a jornada de atendimento, são "
        "recebidos múltiplos tipos de dados sensíveis: cadastros, sintomas clínicos, dados de exames "
        "citológicos e laboratoriais, respostas a formulários, relatos escritos e histórico pregresso. "
        "Essas informações demandam triagem rápida e precisa por parte dos profissionais para priorizar "
        "casos graves, orientar condutas imediatas e consultar protocolos oficiais de referência.\n\n"
        "O desafio proposto para a Fase 5 consiste em desenvolver a Guardiã AI, uma solução tecnológica "
        "robusta capaz de receber dados clínicos e relatos de atendimento, empregando diferentes técnicas "
        "de Inteligência Artificial para apoiar a análise inicial do caso.\n\n"
        "O desafio estabelece a atuação em dois contextos essenciais: saúde da mulher e segurança da mulher.\n\n"
        "Cenário Principal Escolhido pelo Grupo:\n"
        "O grupo selecionou como cenário principal o Apoio à Triagem e Estratificação de Risco de Câncer de Mama "
        "(contexto de Saúde da Mulher), integrando o processamento de relatos clínicos com sanitização e "
        "mascaramento de dados sensíveis (LGPD) e identificação de sinais de alerta (contexto de Segurança da Mulher). "
        "A solução combina Machine Learning com calibração de limiar seguro focado em Recall, explicabilidade diagnóstica "
        "por SHAP, consulta automatizada a protocolos clínicos oficiais (RAG), síntese em linguagem natural via LLM, "
        "guardrails de segurança e homologação obrigatória da decisão médica pelo profissional responsável."
    )
    
    # 2. P[35] - OBJETIVO
    doc.paragraphs[35].text = (
        "Desenvolver e disponibilizar um produto completo de Inteligência Artificial para suporte à triagem "
        "clínica de câncer de mama, demonstrando uma jornada ponta a ponta funcional, segura e responsável.\n\n"
        "Fluxo Operacional da Solução:\n"
        "A aplicação implementa o fluxo completo exigido pelo desafio:\n"
        "Entrada dos dados (citologia PAAF e relato) → Validação e sanitização (LGPD e injeção de prompt) → "
        "Predição por Machine Learning → Interpretabilidade com SHAP → Extração de achados e sinais de alerta → "
        "Consulta a protocolos oficiais (RAG) → Interpretação contextualizada por LLM → Guardrails de saída → "
        "Apresentação na interface web → Homologação da Decisão Médica Final (Human-in-the-Loop) → "
        "Persistência em Trilha de Auditoria Encadeada com SHA-256.\n\n"
        "Diretrizes Mandatórias de Governança e Ética Médica:\n"
        "• \"A decisão final sempre deverá permanecer com uma pessoa responsável pelo atendimento.\"\n"
        "• A Guardiã AI funciona estritamente como um sistema de apoio ao profissional, e não como substituição da decisão humana.\n"
        "• A aplicação não realiza diagnóstico médico definitivo, não prescreve medicamentos e não toma decisões clínicas de forma automática.\n"
        "• Todas as ações, estimativas, consultas ao RAG e decisões médicas são imutavelmente registradas em trilha de auditoria encadeada (SHA-256)."
    )
    
    # 3. P[44] - ENTREGÁVEIS: Relatório técnico com:
    doc.paragraphs[44].text = (
        "● Problema Escolhido e Escopo Clínico: Justificativa do cenário de triagem de câncer de mama, contexto epidemiológico e fronteiras éticas do produto (design for safety);\n"
        "● Conjunto de Dados e Análise Exploratória: Exploração detalhada do dataset público e anonimizado UCI WDBC (Wisconsin Diagnostic Breast Cancer) com 569 casos e 30 atributos morfológicos celulares nucleares de punção aspirativa por agulha fina (PAAF);\n"
        "● Preparação dos Dados e Prevenção de Vazamento: Divisão estratificada 80/20 (455 treino / 114 teste), padronização ajustada exclusivamente no conjunto de treino e validação cruzada estratificada de 5 folds;\n"
        "● Modelos de Machine Learning e Comparativo: Treinamento, otimização (GridSearchCV) e comparação de 4 algoritmos: Regressão Logística, Random Forest, Gradient Boosting e XGBoost;\n"
        "● Métricas e Impacto Clínico dos Erros: Avaliação via Accuracy, Precision, Recall, F1-Score, ROC-AUC, PR-AUC e Brier Score, com calibração de limiar seguro de operação (t_low = 0,09) priorizando alto Recall para mitigar falsos negativos;\n"
        "● Interpretabilidade e Explicabilidade com SHAP: Aplicação de valores de Shapley locais e globais, detalhando as variáveis celulares que mais aumentam ou reduzem a suspeita clínica;\n"
        "● Integração com LLM e Guardrails: Uso de modelo de linguagem (Claude 3.5 Sonnet / Anthropic com fallback determinístico offline) para síntese clara em linguagem natural, com bloqueio estrito de prescrições e diagnósticos definitivos;\n"
        "● RAG (Retrieval-Augmented Generation): Base de conhecimento com 7 documentos oficiais (INCA, Ministério da Saúde, BI-RADS e fluxos SUS), com recuperação por TF-IDF e cosseno (Hit@1 = 100%, Hit@3 = 100%);\n"
        "● Arquitetura e Orquestração: Grafo de estados implementado com LangGraph (8 etapas), interface web em Streamlit e API RESTful em FastAPI;\n"
        "● Governança, LGPD e Auditoria: Mascaramento de dados pessoais (CPF, e-mail, telefone), proteção contra injeção de prompt, controle de acesso RBAC e trilha de auditoria encadeada com hash SHA-256;\n"
        "● Limitações e Responsabilidade Médica: Análise das limitações amostrais e reafirmação do compromisso de que a decisão final sempre deverá permanecer com o profissional de saúde responsável."
    )
    
    # 4. P[48] - ENTREGÁVEIS: Vídeo demonstrando:
    doc.paragraphs[48].text = (
        "● Apresentação conceitual do problema e justificativa clínica da triagem de câncer de mama na saúde da mulher;\n"
        "● Demonstração prática da aplicação funcionando de ponta a ponta na interface web interativa (Streamlit);\n"
        "● Jornada completa de atendimento: login com controle de acesso RBAC, inserção de medidas citológicas reais (WDBC) e relato clínico livre;\n"
        "● Mecanismos de segurança em tempo real: mascaramento automático de dados sensíveis (CPF, e-mail, telefone) e detecção de sinais de alerta (red flags);\n"
        "● Execução do modelo de Machine Learning, calibração do limiar operacional e explicabilidade visual por valores SHAP;\n"
        "● Consulta interativa ao assistente RAG sobre protocolos oficiais de conduta (classificação BI-RADS e fluxos SUS), com citação de fontes;\n"
        "● Geração de síntese explicativa em linguagem natural pela LLM contextualizada com as evidências do caso;\n"
        "● Homologação da Decisão Médica Final pelo profissional responsável, com inclusão de CRM, conduta clínica e termo de responsabilidade soberana;\n"
        "● Módulo de governança: visualização do histórico e comprovação da integridade da trilha de auditoria encadeada (SHA-256 estilo blockchain);\n"
        "● Demonstração da API REST (FastAPI) e execução da suíte de testes automatizados (pytest com 14 testes aprovados)."
    )
    
    # 5. P[60] - RELATÓRIO TÉCNICO: Título
    doc.paragraphs[60].text = "## Sistema Guardiã AI — Triagem e Estratificação de Risco de Câncer de Mama"
    
    # 6. P[62] - RELATÓRIO TÉCNICO: Descrição introdutória
    doc.paragraphs[62].text = (
        "Este relatório técnico detalha a especificação da solução de Inteligência Artificial Guardiã AI, "
        "desenvolvida para apoio a profissionais de saúde na triagem e estratificação de risco de câncer de mama. "
        "A plataforma combina dados citológicos reais (UCI WDBC), modelos de Machine Learning otimizados com foco em "
        "sensibilidade (Recall), explicabilidade clínica por SHAP, Large Language Model (LLM), recuperação de diretrizes "
        "clínicas oficiais por RAG, governança robusta com trilha de auditoria encadeada e a garantia de que a decisão "
        "final sempre deverá permanecer com uma pessoa responsável pelo atendimento."
    )
    
    # 7. P[66] - RELATÓRIO TÉCNICO: Conteúdo detalhado das seções
    # Substituiremos o parágrafo 66 por múltiplos parágrafos estruturados do relatório técnico
    p66 = doc.paragraphs[66]
    
    report_sections = [
        ("1. Resumo Executivo e Problema Clínico", "Heading 3", [
            "O câncer de mama é a neoplasia maligna mais frequente e a principal causa de mortalidade por câncer entre mulheres no Brasil e no mundo. O prognóstico está diretamente associado à precocidade do diagnóstico: tumores diagnosticados em estágios iniciais alcançam taxas de cura superiores a 90%, ao passo que atrasos no encaminhamento clínico elevam os custos assistenciais e reduzem severamente a sobrevida das pacientes.",
            "Nos serviços de acolhimento e triagem, os profissionais de saúde enfrentam o desafio de analisar simultaneamente relatos clínicos de anamnese, histórico de fatores de risco e laudos morfológicos de exames de biópsia por punção aspirativa por agulha fina (PAAF). A Guardiã AI foi desenvolvida precisamente para atuar como ferramenta de apoio à decisão nesse ponto crítico do fluxo assistencial.",
            "Fronteiras de Segurança do Produto (Design for Safety): A Guardiã AI opera sob premissas rigorosas: (i) a IA nunca fecha um caso de forma autônoma; (ii) o sistema não emite diagnóstico definitivo nem prescreve medicamentos; (iii) a decisão clínica final e a conduta terapêutica permanecem obrigatoriamente sob a responsabilidade exclusiva do profissional de atendimento."
        ]),
        ("2. Conjunto de Dados (UCI WDBC) e Análise Exploratória", "Heading 3", [
            "A modelagem preditiva utiliza o dataset público e anonimizado Wisconsin Diagnostic Breast Cancer (WDBC), disponibilizado pelo repositório UCI Machine Learning. O dataset compreende 569 casos clínicos, sendo 357 benignos (62,7%) e 212 malignos (37,3%).",
            "Cada registro é composto por 30 variáveis contínuas que descrevem as propriedades geométricas e de textura dos núcleos celulares observados em lâminas de PAAF digitalizadas. Os atributos derivam de 10 características morfológicas fundamentais: raio, textura, perímetro, área, suavidade, compacidade, concavidade, pontos côncavos, simetria e dimensão fractal.",
            "Para cada núcleo celular, são calculadas três medidas estatísticas: a média (mean), o erro padrão (standard error - se) e o pior valor observado (worst), totalizando os 30 atributos. A análise exploratória evidenciou forte correlação positiva com malignidade nas variáveis associadas à dimensão celular (worst perimeter, worst radius, worst area) e às reentrâncias da membrana nuclear (worst concave points)."
        ]),
        ("3. Preparação dos Dados e Prevenção de Vazamento (Data Leakage)", "Heading 3", [
            "Para garantir validade metodológica e prevenir vazamento de dados (data leakage), o dataset foi dividido em 80% para treinamento (455 amostras) e 20% para teste final cego (114 amostras), com estratificação pela classe de malignidade e semente aleatória fixa (seed=42).",
            "O conjunto de teste cego foi mantido estritamente isolado durante todas as etapas de calibração, seleção de modelos e ajuste de hiperparâmetros. Todas as transformações de pré-processamento (padronização StandardScaler e imputação de valores faltantes) foram ajustadas exclusivamente nos dados de treino e apenas aplicadas ao conjunto de teste.",
            "A validação interna dos modelos utilizou validação cruzada estratificada de 5 folds (StratifiedKFold) com otimização exaustiva de hiperparâmetros via GridSearchCV, assegurando estimativas de desempenho robustas e reprodutíveis."
        ]),
        ("4. Modelos de Machine Learning e Comparativo de Desempenho", "Heading 3", [
            "Foram treinados e comparados quatro algoritmos representativos de diferentes famílias de aprendizado supervisionado: Regressão Logística (linear calibrada), Random Forest (ensemble bagging), Gradient Boosting (boosting sequencial) e XGBoost (gradient boosting com regularização de segunda ordem).",
            "Resultados comparativos na validação cruzada (CV 5-fold, média ± desvio-padrão):\n"
            "• Regressão Logística: Accuracy = 97,4% | Recall = 95,3% | F1-Score = 96,4% | ROC-AUC = 0,995\n"
            "• Gradient Boosting: Accuracy = 97,4% | Recall = 95,9% | F1-Score = 96,4% | ROC-AUC = 0,993\n"
            "• XGBoost: Accuracy = 97,1% | Recall = 95,3% | F1-Score = 96,1% | ROC-AUC = 0,994\n"
            "• Random Forest: Accuracy = 96,0% | Recall = 93,5% | F1-Score = 94,6% | ROC-AUC = 0,989",
            "Modelo Selecionado para Produção: A Regressão Logística (penalização L2, C=0.1, solver lbfgs) foi escolhida como o modelo principal da solução. As razões clínicas e técnicas fundamentais para esta escolha incluem: excelente calibração intrínseca de probabilidades (Brier Score = 0,021), comportamento estritamente monotônico e convexo (previsibilidade clínica), alta interpretabilidade e empate técnico no topo com os modelos baseados em árvores, sem o risco de overfitting."
        ]),
        ("5. Calibração de Limiares e Análise de Impacto dos Erros Clínicos", "Heading 3", [
            "Em problemas de triagem oncológica, os custos e impactos dos erros de classificação são profundamente assimétricos. Um Falso Positivo (FP) gera ansiedade temporária e encaminhamento para exames complementares (mamografia diagnóstica ou ultrassom), com risco clínico mitigado. Já um Falso Negativo (FN) significa liberar inadvertidamente uma paciente com neoplasia ativa, postergando o tratamento e reduzindo drasticamente a sobrevida.",
            "Por essa razão, a Guardiã AI rejeita o limiar arbitrário de 0,5 e adota uma política de estratificação em três faixas de risco com dois limiares calibrados:\n"
            "• Limiar de Segurança (t_low = 0,09): Casos com probabilidade estimada < 0,09 são classificados como 'Baixo Risco' (rastreamento de rotina).\n"
            "• Faixa Intermediária (0,09 ≤ P < 0,43): Casos classificados como 'Risco Intermediário', demandando reavaliação ou exames complementares.\n"
            "• Limiar de Alto Risco (t_high = 0,43): Casos com probabilidade ≥ 0,43 são categorizados como 'Alto Risco', priorizados imediatamente para biópsia e mastologia.",
            "Desempenho no Teste Cego (114 casos não vistos): Aplicado o limiar t_low = 0,09, o modelo obteve:\n"
            "Recall = 97,6% (41 de 42 casos malignos detectados com segurança; apenas 1 FN);\n"
            "Precisão = 85,4% | Acurácia = 93,0% | F1-Score = 91,1% | ROC-AUC = 0,996 | PR-AUC = 0,994 | Brier Score = 0,021."
        ]),
        ("6. Interpretabilidade e Explicabilidade Clínica com SHAP", "Heading 3", [
            "Para assegurar transparência e confiança por parte da equipe médica, a predição é acompanhada de análise de explicabilidade baseada em SHAP (SHapley Additive exPlanations), fundamentada na teoria dos jogos cooperativos.",
            "Para cada caso analisado, o sistema calcula os valores SHAP exatos e renderiza um gráfico de barras horizontal destacando as variáveis que mais influenciaram a predição (em vermelho para aumento de suspeita e azul para redução de suspeita).",
            "A análise global confirma que as variáveis de maior peso clínico são: pior número de pontos côncavos (worst concave points), pior raio (worst radius), pior perímetro (worst perimeter), textura média (mean texture) e raio médio (mean radius). O sistema inclui uma legenda explicativa informando que variáveis morfologicamente correlacionadas dividem a contribuição SHAP."
        ]),
        ("7. Mecanismo de Recuperação Baseado em Protocolos (RAG)", "Heading 3", [
            "A Guardiã AI incorpora um motor RAG (Retrieval-Augmented Generation) com base de conhecimento especializada armazenada em Markdown (data/kb/), contemplando 7 diretrizes oficiais:\n"
            "1. 01_rastreamento_mama.md: Diretrizes do Ministério da Saúde e INCA para rastreamento mamográfico bienal (50 a 69 anos);\n"
            "2. 02_sinais_alerta.md: Sinais e sintomas de alerta (nódulo palpável, descarga papilar sanguinolenta, retração cutânea);\n"
            "3. 03_birads.md: Sistema de laudo e condutas da classificação BI-RADS (categorias 0 a 6);\n"
            "4. 04_citologia_wdbc.md: Interpretação de dados de PAAF e citologia oncótica;\n"
            "5. 05_limites_ia_governanca.md: Diretrizes de bioética, governança clínica e limites de atuação da IA;\n"
            "6. 06_fluxo_encaminhamento.md: Linhas de cuidado e regulação de encaminhamentos no SUS;\n"
            "7. 07_fatores_risco.md: Fatores de risco genéticos (mutações BRCA), idade e densidade mamária.",
            "Os documentos foram segmentados em 29 trechos temáticos indexados por representação vetorial TF-IDF e métrica de similaridade por cosseno. O motor recupera os trechos mais relevantes (k=3), retornando obrigatoriamente o nome do arquivo, a seção e o escore de similaridade.",
            "Avaliação do RAG: Foram submetidas 19 consultas clínicas de validação e 3 consultas fora do domínio (out-of-domain). O RAG alcançou Hit@1 = 100% e Hit@3 = 100%, rejeitando com sucesso consultas irrelevantes (limiar mínimo de similaridade = 0,08)."
        ]),
        ("8. Integração com LLM e Guardrails de Segurança", "Heading 3", [
            "A síntese contextualizada é gerada por Large Language Model (Claude 3.5 Sonnet via API Anthropic), utilizando um prompt estruturado contendo quatro blocos delimitados: dados do caso, predição do modelo, sinais de alerta extraídos do relato e trechos recuperados da base RAG.",
            "Modo Offline e Resiliência: Para garantir disponibilidade ininterrupta mesmo em cenários de instabilidade de conexão ou ausência de chave de API, o sistema conta com um motor determinístico offline que sintetiza as mesmas evidências clínicas em texto estruturado por regras, servindo como fallback transparente.",
            "Guardrails de Entrada (LGPD e Sanitização): Expressões regulares identificam e mascaram automaticamente dados pessoais sensíveis (CPF, e-mail e telefone) presentes no relato livre. Além disso, um classificador baseado em padrões heurísticos detecta e bloqueia tentativas de injeção de prompt.",
            "Guardrails de Saída: O texto gerado é filtrado por expressões regulares que bloqueiam terminantemente afirmações de diagnóstico definitivo ('você tem câncer'), prescrições farmacológicas ('tome tamoxifeno 20mg') ou exclusões sumárias de doença. Caso uma infração seja detectada, o texto é substituído por mensagem segura e o evento é registrado na auditoria com anexação compulsória do disclaimer de responsabilidade médica."
        ]),
        ("9. Arquitetura e Orquestração do Pipeline com LangGraph", "Heading 3", [
            "O pipeline principal da aplicação é orquestrado como um grafo de estados direcionado compilado com a biblioteca LangGraph, composto por 8 etapas coordenadas:\n"
            "1. validate: Validação de tipos, intervalo dos 30 atributos citológicos, sanitização do relato livre e mascaramento de PII;\n"
            "2. model: Execução do modelo de Machine Learning, cálculo de probabilidade e classificação na faixa de risco operacional;\n"
            "3. findings: Extração de sinais de alerta clínicos (red flags) do relato (ex.: nódulo palpável, retração de pele);\n"
            "4. retrieve: Consulta vetorial à base de conhecimento (RAG) buscando protocolos correspondentes aos achados;\n"
            "5. generate: Geração da análise contextualizada em linguagem natural via LLM (ou fallback offline);\n"
            "6. guard: Verificação de guardrails de saída, conformidade e garantia do disclaimer obrigatório;\n"
            "7. persist: Gravação transacional do caso no banco de dados SQLite e inserção na trilha de auditoria encadeada;\n"
            "8. human_review: Homologação da decisão médica final pelo profissional responsável (Human-in-the-Loop)."
        ]),
        ("10. Segurança, Controle de Acesso (RBAC) e Trilha de Auditoria Encadeada", "Heading 3", [
            "Controle de Acesso Baseado em Papéis (RBAC): O sistema define três perfis com permissões estritas:\n"
            "• admin: Criação e gestão de usuários, execução de análises, visualização de histórico geral e auditoria;\n"
            "• profissional: Execução de análises, interação com o assistente RAG, homologação da decisão médica final e visualização do próprio histórico;\n"
            "• auditor: Acesso somente-leitura ao histórico consolidado e verificação de integridade da trilha de auditoria.",
            "Mecanismos Criptográficos de Segurança: Senhas armazenadas com hash bcrypt (12 rounds); bloqueio temporário de conta por 15 minutos após 5 tentativas inválidas; sessões com expiração automática após 30 minutos de inatividade; tokens JWT (algoritmo HS256) na API REST; e pseudonimização de identificadores de casos por HMAC-SHA256.",
            "Trilha de Auditoria com Hash Encadeado (SHA-256): Inspirada na tecnologia blockchain, a tabela de auditoria (audit) encadeia criptograficamente cada registro (h_i = SHA-256(h_{i-1} | timestamp | ator | ação | detalhe)). Caso qualquer registro anterior seja adulterado diretamente no banco de dados, a função verify_audit_chain() detecta a inconsistência e aponta exatamente o primeiro registro violado."
        ]),
        ("11. Homologação da Decisão Médica Final (Human-in-the-Loop)", "Heading 3", [
            "Em estrito cumprimento à diretriz mandatória de que 'a decisão final sempre deverá permanecer com uma pessoa responsável pelo atendimento', o fluxo da Guardiã AI culmina obrigatoriamente na homologação médica.",
            "Após a exibição da estimativa de risco, gráfico SHAP e protocolos RAG, o profissional dispõe de um formulário estruturado para registrar:\n"
            "• Parecer Clínico: Concordância com alto risco / investigação prioritária; concordância com baixo risco / seguimento de rotina; discordância com reclassificação clínica; solicitação de exames complementares (USG/RNM/mamografia); ou conduta personalizada;\n"
            "• Identificação Técnica: Registro profissional (CRM/COREN);\n"
            "• Conduta Detalhada: Prescrição de exames, orientações ao paciente e justificativa clínica;\n"
            "• Termo de Ciência e Responsabilidade: Checkbox obrigatória declarando que a decisão final permanece sob sua responsabilidade soberana.",
            "A decisão é gravada de forma vinculada ao caso e registrada com evento HUMAN_REVIEW na trilha de auditoria encadeada. Na tela de histórico, o status é exibido como '✅ Concluído pelo médico' ou '⏳ Pendente de decisão médica', permitindo homologação posterior de casos pendentes."
        ]),
        ("12. Ambiente de Execução e Suíte de Testes Automatizados", "Heading 3", [
            "A aplicação foi concebida para execução simples e flexível:\n"
            "• Containerização: Dockerfile multi-estágio e docker-compose.yml orquestrando a interface web (Streamlit na porta 8501) e a API REST (FastAPI na porta 8000), com volumes persistentes para banco e chaves criptográficas;\n"
            "• Execução Local: Ambiente virtual Python 3.12 com requirements.txt e script CLI administrativo para inicialização e gestão de usuários.",
            "Validação Automatizada (pytest): O projeto conta com 14 testes automatizados cobrindo: validação de senhas fracas, proteção de hashes, bloqueio de tentativas sucessivas (lockout), ciclo de tokens JWT, controle RBAC, fluxo completo benigno/maligno, divergência entre modelo e relato clínico, sanitização de PII e injeção de prompt, guardrails de saída, cobertura do RAG, integridade e detecção de adulteração na auditoria, endpoints da API e homologação da decisão médica final. Todos os 14 testes executam com 100% de aprovação."
        ]),
        ("13. Limitações da Solução e Trabalhos Futuros", "Heading 3", [
            "Limitações Reconhecidas:\n"
            "1. Tamanho do Dataset: O conjunto UCI WDBC (569 casos) possui valor didático e histórico consolidado, mas não substitui a diversidade demográfica de bases populacionais contemporâneas;\n"
            "2. Ausência de Validação Clínica Externa: O modelo não foi submetido a ensaio clínico multicêntrico e não pode ser utilizado para diagnóstico hospitalar sem validação regulatória;\n"
            "3. Origem das Variáveis: As 30 medidas citológicas derivam de imagens digitalizadas de biópsia por agulha fina (PAAF), dependendo de operador especializado para aquisição prévia dos atributos celulares;\n"
            "4. Infraestrutura de Persistência: O uso de SQLite local com segredo em arquivo é ideal para demonstração e avaliação, mas deve ser substituído por PostgreSQL e cofre de chaves (Key Vault) em ambiente de produção clínico.",
            "Evoluções Futuras Recomendadas: Integração de visão computacional multimodal para extração direta de características a partir de imagens histopatológicas (conforme explorado na Fase 4); expansão da base de conhecimento RAG com embeddings neurais densos (BioBERT/ClinicalBERT); e integração via padrão HL7/FHIR com prontuários eletrônicos do paciente (PEP)."
        ])
    ]
    
    h2_style = next(s for s in doc.styles if s.style_id == 'Heading2')
    h3_style = next(s for s in doc.styles if s.style_id == 'Heading3')
    body_style = doc.paragraphs[35].style  # Body Text

    # Inserir as seções antes de P[66]
    for title, style_type, paragraphs in report_sections:
        h = p66.insert_paragraph_before(title)
        h.style = h2_style if style_type == "Heading 2" else h3_style
        for text in paragraphs:
            p_body = p66.insert_paragraph_before(text)
            p_body.style = body_style
    
    # Remover o parágrafo 'XXXXX' original do P[66]
    p66_elem = p66._element
    p66_elem.getparent().remove(p66_elem)
    
    # 8. P[70] (ou parágrafo com 'XXXXX' sob CONCLUSÃO)
    # Como removemos P[66], vamos encontrar o parágrafo sob CONCLUSÃO que contenha XXXXX
    for p in doc.paragraphs:
        if "XXXXX" in p.text:
            p.text = (
                "A Guardiã AI cumpre integralmente os objetivos pedagógicos, técnicos e éticos estabelecidos para "
                "o Tech Challenge Fase 5. O projeto demonstra na prática como transitar de modelos analíticos isolados "
                "para uma plataforma integrada de apoio à decisão clínica, unindo com sucesso Dados, Machine Learning, "
                "SHAP, LLM, RAG, Governança e Aplicação.\n\n"
                "A seleção da Regressão Logística calibrada com limiar operacional de segurança focado em Recall "
                "(97,6% no teste cego com apenas 1 falso negativo em 42 casos) assegura sensibilidade clínica máxima na triagem, "
                "enquanto a explicabilidade por SHAP e o embasamento documental do RAG em diretrizes do INCA/MS e BI-RADS "
                "conferem total transparência e suporte técnico ao profissional de saúde.\n\n"
                "Por fim, sob a perspectiva de bioética e conformidade legal, a aplicação estabelece salvaguardas robustas: "
                "sanitização de dados sensíveis conforme a LGPD, guardrails contra alucinações e prescrições indevidas, e "
                "trilha de auditoria encadeada com SHA-256 à prova de adulterações. Acima de tudo, o sistema consagra "
                "a diretriz mandatória de que \"a decisão final sempre deverá permanecer com uma pessoa responsável pelo "
                "atendimento\", garantindo que a Inteligência Artificial permaneça como uma aliada do julgamento clínico humano "
                "em defesa da saúde e da vida da mulher."
            )
            break
            
    doc.save(docx_path)
    print(f"Arquivo salvo com sucesso: {docx_path}")

if __name__ == "__main__":
    for path in [
        "/media/pc/hdd/Sandisk/Cursos/PosTech/TechChallenge05/guardia-ai/Tech Challenge 05.docx",
        "/media/pc/hdd/Sandisk/Cursos/PosTech/TechChallenge05/Tech Challenge 05.docx"
    ]:
        if os.path.exists(path):
            update_document(path)
